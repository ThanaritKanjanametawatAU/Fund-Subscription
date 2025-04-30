import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, numbers
import os
import io
import base64
import sys
import traceback

def process_fund_subscription(master_file_content, data_file_content, report_date):
    """
    Process fund subscription files and generate a report.
    
    Args:
        master_file_content: The content of the MasterFile.xlsx file
        data_file_content: The content of the Data.xlsx file
        report_date: The report date in DD/MM/YY format
        
    Returns:
        The content of the generated Excel file
    """
    try:
        # Validate inputs
        if not master_file_content or len(master_file_content) < 10:
            raise ValueError("Master file content is empty or too small")
        if not data_file_content or len(data_file_content) < 10:
            raise ValueError("Data file content is empty or too small")
        if not report_date or len(report_date) < 5:
            raise ValueError("Report date is invalid or missing")
            
        print(f"Starting to process files. Master file size: {len(master_file_content)} bytes, Data file size: {len(data_file_content)} bytes")
        
        # Read data from file-like objects
        try:
            master_file = pd.read_excel(io.BytesIO(master_file_content))
            print(f"Successfully loaded master file with {len(master_file)} rows")
        except Exception as e:
            print(f"Error reading master file: {str(e)}")
            traceback.print_exc()
            raise ValueError(f"Failed to read master file: {str(e)}")
            
        try:
            data_file = pd.read_excel(io.BytesIO(data_file_content), skiprows=2)
            print(f"Successfully loaded data file with {len(data_file)} rows")
        except Exception as e:
            print(f"Error reading data file: {str(e)}")
            traceback.print_exc()
            raise ValueError(f"Failed to read data file: {str(e)}")
        
        # Clean data
        try:
            data_cleaned = data_file.dropna(axis=1, how='all', subset=[data_file.index[0]])
            total_row_index = data_cleaned[data_cleaned.iloc[:,0] == 'Total by Bank'].index.min()
            if pd.isna(total_row_index):
                print("Warning: 'Total by Bank' row not found, using all data rows")
                total_row_index = len(data_cleaned)
            data_cleaned = data_cleaned.iloc[0: total_row_index]
            
            # Check if first row exists
            if len(data_cleaned) == 0 or not data_cleaned.iloc[0].any():
                raise ValueError("Data file is empty or missing header row")
                
            data_cleaned.columns = data_cleaned.iloc[0]
            
            # Check for required columns
            required_columns = ["Fund / Bank", "Total by Fund"]
            for col in required_columns:
                if col not in data_cleaned.columns:
                    raise ValueError(f"Required column '{col}' not found in data file")
            
            # Drop columns safely
            columns_to_drop = ["Total by Fund", "หมายเหตุ"]
            for col in columns_to_drop:
                if col in data_cleaned.columns:
                    data_cleaned = data_cleaned.drop(columns=[col])
                else:
                    print(f"Warning: Column '{col}' not found for dropping")
                    
            data_cleaned = data_cleaned[1:]
            print(f"Data cleaning completed. Cleaned data has {len(data_cleaned)} rows and {len(data_cleaned.columns)} columns")
        except Exception as e:
            print(f"Error cleaning data: {str(e)}")
            traceback.print_exc()
            raise ValueError(f"Failed to clean data: {str(e)}")
        
        # Create Excel workbook
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            print("Created Excel workbook")
            
            # Apply styles
            center_aligned_text = Alignment(horizontal='center', vertical='center')
            
            sheet.merge_cells('A1:E1')
            sheet.merge_cells('A2:E2')
            sheet['A1'].alignment = center_aligned_text
            sheet['A2'].alignment = center_aligned_text
            sheet["A1"] = "บริษัท หลักทรัพย์จัดการกองทุน แลนด์ แอนด์ เฮ้าส์ จำกัด"
            sheet["A2"] = "LAND AND HOUSES FUND MANAGEMENT CO.,LTD."
            sheet["C4"] = "JV"
            sheet["C5"] = "Date :"
            sheet["C6"] = "BackDate :"
            sheet["D5"] = report_date
            
            cells = ["A", "B", "C", "D"]
            texts = ["Account Code", "PARTICULAR", "DR.", "CR."]
            for i in range(len(cells)):
                sheet[cells[i] + str(9)] = texts[i]
        except Exception as e:
            print(f"Error creating workbook: {str(e)}")
            traceback.print_exc()
            raise ValueError(f"Failed to create Excel workbook: {str(e)}")
        
        # Define borders
        right_border = Border(right=Side(style='thin'))
        all_border = Border(bottom=Side(style='thin'), right=Side(style='thin'), left=Side(style='thin'), top=Side(style='thin'))
        sum_border = Border(bottom=Side(style='double'), right=Side(style='thin'), left=Side(style='thin'), top=Side(style='thin'))
        
        def rightBorder(row):
            for c in ["A","B","C","D"]:
                cell = sheet[c+str(row)]
                cell.border = right_border
        
        # Apply accounting styles
        rightBorder(10)
        
        for c in ["A","B","C","D"]:
            cell = sheet[c+str(9)]
            cell.border = all_border
        
        # Add data entries
        try:
            currentRow = 11
            entries_processed = 0
            
            for col in data_cleaned.columns:
                for index, rowData in data_cleaned.iterrows():
                    if col != "Fund / Bank" and rowData[col] is not np.nan:
                        # Record One Entry
                        try:
                            rowMasterFile = master_file[master_file['Name'] == col]
                            if len(rowMasterFile) == 0:
                                print(f"Warning: No match in master file for '{col}'")
                                continue
                                
                            sheet["A"+str(currentRow)] = str(rowMasterFile["AccountNumber"].reset_index(drop=True)[0])
                            sheet["A"+str(currentRow)].alignment = center_aligned_text
                            
                            sheet["B"+str(currentRow)] = rowMasterFile["BankAccountNumber"].reset_index(drop=True)[0]
                            
                            sheet["C"+str(currentRow)] = rowData[col]
                            sheet["C"+str(currentRow)].number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1
                            
                            rightBorder(currentRow)
                            
                            currentRow += 1
                            sheet["A"+str(currentRow)] = "2090501020"
                            sheet["A"+str(currentRow)].alignment = center_aligned_text
                            
                            sheet["B"+str(currentRow)] = "          FUND SUBSCRIPTION"
                            
                            sheet["D"+str(currentRow)] = rowData[col]
                            sheet["D"+str(currentRow)].number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1
                            
                            rightBorder(currentRow)
                            
                            currentRow += 1
                            sheet["B"+str(currentRow)] = "ลูกค้านำเงินจองซื้อกองทุน/" + rowData["Fund / Bank"] + "  ลว." + report_date
                            
                            rightBorder(currentRow)
                            rightBorder(currentRow+1)
                            
                            currentRow += 2
                            rightBorder(currentRow)
                            
                            entries_processed += 1
                        except Exception as entry_error:
                            print(f"Warning: Error processing entry for {col}, row {index}: {str(entry_error)}")
                            # Continue with next entry
            
            print(f"Processed {entries_processed} entries")
            
            if entries_processed == 0:
                raise ValueError("No valid entries were processed. Check your input files.")
        except Exception as e:
            print(f"Error adding data entries: {str(e)}")
            traceback.print_exc()
            raise ValueError(f"Failed to add data entries: {str(e)}")
        
        # Make Summation
        for i in range(10):
            rightBorder(currentRow + i)
        currentRow += 9
        
        sheet["C"+str(currentRow)] = f"=SUM(C10:C{currentRow-1})"
        sheet["C"+str(currentRow)].number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1
        sheet["D"+str(currentRow)] = f"=SUM(D10:D{currentRow-1})"
        sheet["D"+str(currentRow)].number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1
        
        for c in ["A","B","C","D"]:
            cell = sheet[c+str(currentRow)]
            cell.border = sum_border
        
        # Expand Column
        columnsWidth = {"A": 15, "B":63, "C":16, "D": 16, "E": 15}
        for column in columnsWidth:
            sheet.column_dimensions[column].width = columnsWidth[column]
        
        # Set Font
        bold_font = Font(name="Arial", size=10, bold=True)
        font_style = Font(name="Arial", size=10)
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = font_style
        
        # Set Bold Font
        cells = ["A1", "C4", "C5", "C6", "D4", "D5", "D6", "A9", "B9", "C9", "D9"]
        for c in cells:
            cell = sheet[c]
            cell.font = bold_font
            if c in ["A1", "A9", "B9", "C9", "D9"]:
                cell.alignment = center_aligned_text
        
        # Save the workbook to a bytes buffer
        try:
            output = io.BytesIO()
            workbook.save(output)
            output.seek(0)
            
            result = output.getvalue()
            print(f"Excel report generated successfully. Size: {len(result)} bytes")
            return result
        except Exception as e:
            print(f"Error saving workbook: {str(e)}")
            traceback.print_exc()
            raise ValueError(f"Failed to save Excel workbook: {str(e)}")
    
    except Exception as e:
        # Handle any errors
        print(f"Critical error in process_fund_subscription: {str(e)}")
        traceback.print_exc()
        raise Exception(f"Error generating report: {str(e)}") 